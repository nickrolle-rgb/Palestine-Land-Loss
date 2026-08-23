"""Attacks on health care, from Insecurity Insight's SiND data.

This is the one source in the project that **cannot be mapped**, and the reason
is the publisher's own protection of the people involved: every record carries
`Geo Precision: censored` and no coordinates at all. Not some records — all
4,408 of them.

That is a decision by people closer to the danger than we are, and it is not
ours to work around. There is no geocoding step here, no matching of place
names to the gazetteer, no plotting at admin-area centroids. The data ships as
counts and the UI presents a table.

What survives censorship is still substantial: the date, the governorate-level
area, the reported perpetrator, and per-incident counts of health workers
killed, injured, kidnapped and arrested, and of facilities destroyed and
damaged. Aggregated, that is the medical picture the PRCS facilities layer
cannot give — that layer says where the ambulances were, this says what
happened to the people in them.

Non-negotiable 6 also applies and is easy here: the public release carries no
event descriptions, so there are no individuals to avoid naming.
"""

from __future__ import annotations

import collections
from typing import Any

from ..fetch import download, retrieved_date
from ..schema import Evidence
from ..sources import SOURCES, resource

DOCUMENT_DATE = "2026-08-17"

COUNT_FIELDS = {
    "Health Workers Killed": "workers_killed",
    "Health Workers Injured": "workers_injured",
    "Health Workers Kidnapped": "workers_kidnapped",
    "Health Workers Arrested": "workers_arrested",
    "Number of Attacks on Health Facilities Reporting Destruction": "facilities_destroyed",
    "Number of Attacks on Health Facilities Reporting Damaged": "facilities_damaged",
}


def load_health_attacks() -> tuple[dict[str, Any], dict[str, Any]]:
    import openpyxl  # noqa: PLC0415

    r = resource("health_attacks")
    path = download(r.url, r.filename)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    ix = {h: i for i, h in enumerate(header) if h}

    totals: collections.Counter = collections.Counter()
    by_year: dict[str, collections.Counter] = {}
    by_area: collections.Counter = collections.Counter()
    by_perpetrator: collections.Counter = collections.Counter()
    incidents = 0
    geocoded = 0

    for row in rows:
        incidents += 1
        lat = row[ix["Latitude"]] if "Latitude" in ix else None
        lon = row[ix["Longitude"]] if "Longitude" in ix else None
        if lat not in (None, "") and lon not in (None, ""):
            geocoded += 1

        date = row[ix["Date"]]
        year = str(date)[:4] if date else "unknown"
        bucket = by_year.setdefault(year, collections.Counter())
        bucket["incidents"] += 1

        by_area[str(row[ix["Admin 1"]] or "No Information")] += 1
        by_perpetrator[str(row[ix["Reported Perpetrator Name"]] or "No Information")] += 1

        for column, key in COUNT_FIELDS.items():
            value = row[ix[column]] if column in ix else None
            if isinstance(value, (int, float)):
                totals[key] += int(value)
                bucket[key] += int(value)

    ev = Evidence(
        source_id="insecurity_insight",
        title="State of Palestine — Attacks on Health Care incident data, 2016–2026",
        url=SOURCES["insecurity_insight"].url,
        document_date=DOCUMENT_DATE,
        retrieved=retrieved_date(r.filename),
        note="Published with all coordinates censored by Insecurity Insight. "
             "Counts only; nothing here is plotted.",
    )

    payload = {
        "incidents": incidents,
        "geocoded": geocoded,
        "censored": incidents - geocoded,
        "totals": dict(totals),
        "by_year": {y: dict(c) for y, c in sorted(by_year.items())},
        "by_area": dict(by_area.most_common()),
        "by_perpetrator": dict(by_perpetrator.most_common(8)),
        "licence": "CC BY-SA 4.0 — Insecurity Insight",
        "evidence": [ev.to_dict()],
    }
    stats = {
        "incidents": incidents,
        "censored": incidents - geocoded,
        "workers_killed": totals["workers_killed"],
        "facilities_destroyed": totals["facilities_destroyed"],
    }
    return payload, stats
